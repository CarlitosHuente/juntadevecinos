from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook

from vecinos.models import Familiar, Vecino
from vecinos.rut import normalizar_rut, validar_rut

COLUMNAS_SOCIOS = (
    "RUT",
    "Nombres",
    "Apellido paterno",
    "Apellido materno",
    "Dirección",
    "Fecha nacimiento (DD-MM-AAAA)",
    "Correo",
    "Teléfono",
    "Fecha ingreso (DD-MM-AAAA)",
)

COLUMNAS_FAMILIARES = (
    "RUT socio titular",
    "Nombre",
    "RUT",
    "Fecha nacimiento (DD-MM-AAAA)",
    "Correo",
    "Teléfono",
)

COLUMNAS_FECHA_SOCIOS = (6, 9)  # F, I
COLUMNAS_FECHA_FAMILIARES = (4,)  # D


def _celda(fila, indice) -> str:
    if indice >= len(fila):
        return ""
    valor = fila[indice]
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.date().strftime("%d-%m-%Y")
    if isinstance(valor, date):
        return valor.strftime("%d-%m-%Y")
    return str(valor).strip()


def _fecha(valor: str):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = (valor or "").strip()
    if not texto:
        return None
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(texto[:10], fmt).date()
        except ValueError:
            continue
    return None


def _formatear_fechas(hoja, indices_columna):
    """Texto DD-MM-AAAA + quotePrefix: Excel no convierte ni muestra AAAA-MM-DD."""
    for indice in indices_columna:
        for fila in hoja.iter_rows(min_row=2, max_row=500, min_col=indice, max_col=indice):
            celda = fila[0]
            if isinstance(celda.value, datetime):
                celda.value = celda.value.strftime("%d-%m-%Y")
            elif isinstance(celda.value, date):
                celda.value = celda.value.strftime("%d-%m-%Y")
            celda.number_format = "@"
            celda.quotePrefix = True


def _libro(columnas) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Datos"
    hoja.append(list(columnas))
    if columnas is COLUMNAS_SOCIOS:
        hoja.append(["12.345.678-5", "Juan", "Pérez", "Soto", "Pasaje 1", "12-05-1980", "", "", "01-03-2024"])
        _formatear_fechas(hoja, COLUMNAS_FECHA_SOCIOS)
    else:
        hoja.append(["12.345.678-5", "Lucas Pérez", "11.111.111-1", "20-08-2010", "", ""])
        _formatear_fechas(hoja, COLUMNAS_FECHA_FAMILIARES)
    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def plantilla_socios() -> bytes:
    return _libro(COLUMNAS_SOCIOS)


def plantilla_familiares() -> bytes:
    return _libro(COLUMNAS_FAMILIARES)


def _filas(archivo, columnas):
    libro = load_workbook(archivo, read_only=True, data_only=True)
    hoja = libro.active
    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        return []
    return filas[1:]


def importar_socios(junta, archivo) -> dict:
    creados, omitidos, errores = 0, 0, []
    vistos = set()
    for i, fila in enumerate(_filas(archivo, COLUMNAS_SOCIOS), start=2):
        rut_crudo = _celda(fila, 0)
        nombres = _celda(fila, 1)
        apellido = _celda(fila, 2)
        if not any([rut_crudo, nombres, apellido]):
            continue
        if not validar_rut(rut_crudo):
            errores.append(f"Fila {i}: RUT inválido ({rut_crudo or 'vacío'}).")
            continue
        rut = normalizar_rut(rut_crudo)
        if rut in vistos or Vecino.objects.filter(junta=junta, rut=rut).exists():
            omitidos += 1
            errores.append(f"Fila {i}: el RUT {rut} ya está registrado.")
            continue
        if not nombres or not apellido:
            errores.append(f"Fila {i}: faltan nombres o apellido paterno.")
            continue
        direccion = _celda(fila, 4)
        if not direccion:
            errores.append(f"Fila {i}: la dirección es obligatoria.")
            continue
        Vecino.objects.create(
            junta=junta,
            rut=rut,
            nombres=nombres,
            apellido_paterno=apellido,
            apellido_materno=_celda(fila, 3),
            direccion=direccion,
            fecha_nacimiento=_fecha(_celda(fila, 5)),
            email=_celda(fila, 6),
            telefono=_celda(fila, 7),
            fecha_ingreso=_fecha(_celda(fila, 8)) or date.today(),
        )
        vistos.add(rut)
        creados += 1
    return {"creados": creados, "omitidos": omitidos, "errores": errores}


def importar_familiares(junta, archivo) -> dict:
    creados, omitidos, errores = 0, 0, []
    vistos = set()
    for i, fila in enumerate(_filas(archivo, COLUMNAS_FAMILIARES), start=2):
        rut_socio = _celda(fila, 0)
        nombre = _celda(fila, 1)
        rut_crudo = _celda(fila, 2)
        if not any([rut_socio, nombre, rut_crudo]):
            continue
        if not validar_rut(rut_socio):
            errores.append(f"Fila {i}: RUT del socio titular inválido.")
            continue
        socio = Vecino.objects.filter(junta=junta, rut=normalizar_rut(rut_socio)).first()
        if not socio:
            errores.append(f"Fila {i}: no hay un socio titular con RUT {normalizar_rut(rut_socio)}.")
            continue
        if not nombre:
            errores.append(f"Fila {i}: el nombre del familiar es obligatorio.")
            continue
        rut = ""
        if rut_crudo:
            if not validar_rut(rut_crudo):
                errores.append(f"Fila {i}: RUT del familiar inválido ({rut_crudo}).")
                continue
            rut = normalizar_rut(rut_crudo)
            if (
                rut in vistos
                or Vecino.objects.filter(junta=junta, rut=rut).exists()
                or Familiar.objects.filter(socio__junta=junta, rut=rut).exists()
            ):
                omitidos += 1
                errores.append(f"Fila {i}: el RUT {rut} ya está en la nómina.")
                continue
            vistos.add(rut)
        Familiar.objects.create(
            socio=socio,
            nombre=nombre,
            rut=rut,
            fecha_nacimiento=_fecha(_celda(fila, 3)),
            email=_celda(fila, 4),
            telefono=_celda(fila, 5),
        )
        creados += 1
    return {"creados": creados, "omitidos": omitidos, "errores": errores}
